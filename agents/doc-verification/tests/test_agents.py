import json
import zipfile

import pytest
from docx import Document as Docx
from PIL import Image
from reportlab.pdfgen import canvas

from dva import clauses, gate, journal, norms, refs
from dva.extract import read


def make_docx(path, lines):
    d = Docx()
    for line in lines:
        d.add_paragraph(line)
    d.save(path)
    return path


def make_scan_pdf(path):
    Image.new("RGB", (600, 800), "white").save(path, "PDF")
    return path


def make_text_pdf(path, pages_with_text, blank_pages=0):
    c = canvas.Canvas(str(path))
    for i in range(pages_with_text):
        c.drawString(72, 720, f"Section {i + 1}. The contractor shall prepare the design documentation in full scope.")
        c.showPage()
    for _ in range(blank_pages):
        c.showPage()
    c.save()
    return path


# ---------- чтение и светофор ----------

def test_gate_green_when_everything_read(tmp_path):
    make_docx(tmp_path / "tz.docx", ["1. Разработать проект."])
    make_text_pdf(tmp_path / "annex.pdf", 2)
    v = gate.check(tmp_path)
    assert v.light == "GREEN", v.report()


def test_gate_red_on_scan_without_text_layer(tmp_path):
    make_docx(tmp_path / "tz.docx", ["1. Разработать проект."])
    make_scan_pdf(tmp_path / "smeta.pdf")
    v = gate.check(tmp_path)
    assert v.light == "RED"
    assert any("скан" in r for d in v.problems for r in d.reasons)


def test_gate_yellow_on_partly_scanned_pdf(tmp_path):
    make_text_pdf(tmp_path / "tz.pdf", 2, blank_pages=1)
    assert gate.check(tmp_path).light == "YELLOW"


def test_gate_red_on_empty_folder(tmp_path):
    assert gate.check(tmp_path).light == "RED"


@pytest.mark.parametrize("name,content", [("empty.docx", b""), ("table.xlsx", b"PK\x03\x04"), ("broken.docx", b"not a zip")])
def test_unreadable_files_are_reported_not_skipped(tmp_path, name, content):
    (tmp_path / name).write_bytes(content)
    doc = read(tmp_path / name)
    assert doc.status == "unread" and doc.reasons


def test_docx_with_entities_is_rejected(tmp_path):
    """XXE / «billion laughs»: в настоящем Word нет DTD — такой файл не разбираем."""
    p = tmp_path / "evil.docx"
    xml = (b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]>'
           b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           b'<w:body><w:p><w:r><w:t>&b;</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("word/document.xml", xml)
    doc = read(p)
    assert doc.status == "unread"
    assert "небезопасный" in doc.reasons[0]


def test_docx_zip_bomb_is_rejected(tmp_path, monkeypatch):
    import dva.extract as ex

    monkeypatch.setattr(ex, "MAX_PART_BYTES", 1000)
    p = tmp_path / "bomb.docx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"<w:document/>" + b" " * 5000)
    assert read(p).status == "unread"


def test_docx_many_parts_hit_total_limit(tmp_path, monkeypatch):
    """Каждая часть под пределом, но колонтитулов сотни — срабатывает общий предел."""
    import dva.extract as ex

    monkeypatch.setattr(ex, "MAX_TOTAL_BYTES", 3000)
    p = tmp_path / "many.docx"
    part = b"<w:hdr xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>" + b" " * 900
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>")
        for i in range(10):
            z.writestr(f"word/header{i}.xml", part)
    doc = read(p)
    assert doc.status == "unread" and "zip-бомба" in doc.reasons[0]


def test_docx_tables_and_footnotes_are_read(tmp_path):
    d = Docx()
    d.add_paragraph("1. Основной текст.")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Объём 2,4 км"
    d.add_paragraph("2. Следующий пункт.")
    d.save(tmp_path / "t.docx")
    text = read(tmp_path / "t.docx").text
    assert "| Объём 2,4 км |" in text
    # строка таблицы не приклеивается к пункту 1 и не рвёт нумерацию
    cl = clauses.split(text)
    assert [(c.number, c.text) for c in cl] == [("1", "Основной текст."), ("2", "Следующий пункт.")]


def test_citation_context_stays_in_its_paragraph():
    (c,) = refs.find_citations("1. Первый пункт без ссылок, довольно длинный текст.\n2. По пункту 5.2.1 ГОСТ Р 90001-2024.")
    assert c.context.startswith("2. По пункту")


# ---------- ссылки на нормативы ----------

NORM = """ГОСТ Р 90001-2024 (учебный)
5.2.1 Проект разрабатывают на весь участок.
5.2.10 Другой пункт.
Таблица 3 — Ведомость
"""


def test_refs_ok_missing_and_unverifiable():
    text = ("Выполнить по пункту 5.2.1 ГОСТ Р 90001-2024. Ведомость по таблице 3 ГОСТ Р 90001-2024. "
            "Знаки по пункту 6.4.9 ГОСТ Р 90001-2024. Учёт по пункту 4.3 ОДМ 90003-2022.")
    lib = {refs.normalize("ГОСТ Р 90001-2024"): NORM}
    result = {(c.number, c.status) for c in refs.check(text, lib)}
    assert ("5.2.1", "ok") in result
    assert ("3", "ok") in result
    assert ("6.4.9", "missing") in result
    assert ("4.3", "unverifiable") in result


def test_clause_number_is_not_matched_as_prefix():
    """«5.2.1» не должен находиться внутри «5.2.10»."""
    assert refs.clause_exists("5.2.10 Другой пункт.", "clause", "5.2.1") is False


def test_designation_does_not_swallow_sentence_dot():
    (c,) = refs.find_citations("Оформить по разделу 7 СП 90002.2023.")
    assert c.document == "СП 90002.2023"


# ---------- аудит библиотеки ----------

def test_norm_audit_finds_substitution_outdated_absent(tmp_path):
    lib = tmp_path / "lib"
    lib.mkdir()
    (lib / "a.txt").write_text("ГОСТ Р 90001-2024\nОрганизация дорожного движения\nИзменение № 1", encoding="utf-8")
    (lib / "b.txt").write_text("СП 90002.2023\nЗдания жилые", encoding="utf-8")
    (lib / "c.txt").write_text("ОДМ 90003-2022\nУчёт интенсивности движения", encoding="utf-8")
    registry = tmp_path / "r.json"
    registry.write_text(json.dumps([
        {"designation": "ГОСТ Р 90001-2024", "title_keywords": ["организация дорожного"], "edition_markers": ["Изменение № 1"]},
        {"designation": "СП 90002.2023", "title_keywords": ["дорожного движения"]},
        {"designation": "ОДМ 90003-2022", "title_keywords": ["интенсивности"], "edition_markers": ["Изменение № 2"]},
        {"designation": "ГОСТ Р 90004-2025", "title_keywords": ["разметка"]},
    ], ensure_ascii=False), encoding="utf-8")
    got = [s.status for s in norms.audit(registry, lib)]
    assert got == ["ok", "substitution", "outdated", "absent"]


# ---------- разметка ТЗ ----------

def test_clauses_split_and_classify(tmp_path):
    rules_path = tmp_path / "rules.json"
    rules_path.write_text(json.dumps({"ЧУЖОЕ": ["лаборатори", "\\bСРО\\b"], "НАШЕ": ["разработать"]}, ensure_ascii=False), encoding="utf-8")
    text = "1. Разработать проект.\n2. Обследование лабораторией.\n3. Соблюдать срок 45 дней.\n45 дней\n4.1 Прочее требование"
    cl = clauses.classify(clauses.split(text), clauses.load_rules(rules_path))
    assert [c.number for c in cl] == ["1", "2", "3", "4.1"]  # «45 дней» из таблицы — не пункт
    assert [c.bucket for c in cl] == ["НАШЕ", "ЧУЖОЕ", "ПРОЧЕЕ", "ПРОЧЕЕ"]  # «срок» не путается с «СРО»


# ---------- журнал и ревизор ----------

def test_journal_append_only_and_review(tmp_path):
    j = journal.Journal(tmp_path / "j.jsonl")
    j.add(journal.Entry("A", "seg", "take", "наш профиль", forecast=10.0))
    j.add(journal.Entry("B", "seg", "take", "наш профиль", forecast=12.0))
    j.add(journal.Entry("C", "seg", "skip", "полевые"))
    with pytest.raises(ValueError):
        j.add(journal.Entry("A", "seg", "take", "дубль"))
    j.outcome("A", "lost", 18.0)
    j.outcome("B", "lost", 20.0)
    text = journal.review(j.entries())
    assert "без исхода 1" in text and "C" in text
    assert "смещение +8.0" in text and "ПРЕДЛОЖЕНИЕ" in text
    # журнал только дописывается: строк больше, чем решений
    assert len((tmp_path / "j.jsonl").read_text(encoding="utf-8").splitlines()) == 5
